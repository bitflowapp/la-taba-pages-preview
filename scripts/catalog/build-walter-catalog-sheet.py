#!/usr/bin/env python3
"""Arma la planilla comercial de La Taba para Walter con los 72 productos recuperados.

    python -I scripts/catalog/build-walter-catalog-sheet.py \
        --backup "$env:TABA_WALTER_CATALOG_BACKUP" \
        --out "$env:TABA_WALTER_CATALOG_XLSX"

Reglas:
  * Del respaldo se lee SOLO data.products. Ninguna otra clave se imprime ni se copia.
  * No se inventan costos, precios actuales ni stock: esas celdas quedan vacías.
    Los valores históricos van rotulados «histórico ago–sep 2026 (vencido)» y las
    sugerencias (categoría, tamaño leído del nombre/código/descripción) como
    «sugerido — confirmar».
  * El vocabulario de categorías sale de js/core/store-taxonomy.js y se cruza con el
    CHECK de 20260828120000_catalogo_multirubro.sql.
  * La planilla y su resumen JSON se escriben FUERA del repositorio (es público).
  * Después de guardar, la planilla se vuelve a abrir y se verifica su contrato.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

EXPECTED_SOURCE_SHA256 = '6983957294de53c7cb3c59ba7aefcc605d2bfb2673b0c5288b6568322b231219'
EXPECTED_ROWS = 72
PRODUCT_KEYS = ['id', 'sku', 'name', 'brand', 'presentation', 'description', 'category', 'price', 'price_status',
                'stock', 'available', 'is_active', 'is_verified', 'is_alcoholic', 'sort_order', 'catalog_origin',
                'created_at', 'updated_at']
SHEETS = ['Instrucciones', 'Productos', 'Inconsistencias', 'Categorías', 'Resumen']
HEADER_ROW, FIRST_ROW = 3, 4
HISTORIC = 'histórico ago–sep 2026 (vencido)'
SUGGESTED = 'sugerido — confirmar'

# letra, título EXACTO (el importador del Panel depende de estos textos), rol, ancho
COLUMNS = [
    ('A', 'SKU', 'reference', 44),
    ('B', 'Producto', 'reference', 26),
    ('C', 'Marca', 'reference', 15),
    ('D', 'Presentación registrada', 'reference', 21),
    ('E', 'Tamaño detectado en el nombre (sugerido — confirmar)', 'suggestion', 25),
    ('F', 'Categoría registrada', 'reference', 15),
    ('G', 'Categoría sugerida (confirmar)', 'suggestion', 18),
    ('H', '¿Alcohol?', 'reference', 9),
    ('I', 'Precio histórico ago–sep 2026 (vencido)', 'reference', 16),
    ('J', 'Stock histórico', 'reference', 9),
    ('K', '¿Estaba disponible?', 'reference', 11),
    ('L', 'Inconsistencias detectadas', 'reference', 42),
    ('M', 'Costo (ARS)', 'input', 14),
    ('N', 'Precio de venta (ARS)', 'input', 14),
    ('O', 'Precio regular antes de la promo (ARS)', 'input', 16),
    ('P', 'Stock actual', 'input', 10),
    ('Q', '¿Disponible para vender?', 'input', 12),
    ('R', 'Ganancia por unidad (ARS)', 'formula', 15),
    ('S', 'Margen sobre venta', 'formula', 11),
    ('T', 'Recargo sobre costo', 'formula', 11),
    ('U', 'Descuento de la promo', 'formula', 11),
    ('V', 'Estado', 'formula', 25),
    ('W', 'Observaciones', 'input', 34),
]
INPUT_LETTERS = [c[0] for c in COLUMNS if c[2] == 'input']
FORMULAS = {
    'R': '=IF(AND(ISNUMBER(N{r}),ISNUMBER(M{r})),N{r}-M{r},"")',
    'S': '=IF(AND(ISNUMBER(N{r}),N{r}>0,ISNUMBER(M{r})),(N{r}-M{r})/N{r},"")',
    'T': '=IF(AND(ISNUMBER(N{r}),ISNUMBER(M{r}),M{r}>0),(N{r}-M{r})/M{r},"")',
    'U': '=IF(AND(ISNUMBER(O{r}),ISNUMBER(N{r}),O{r}>N{r}),(O{r}-N{r})/O{r},"")',
    # INT va en su propio IF: OR() de Excel propaga el error de INT(texto).
    'V': ('=IF(OR(NOT(ISNUMBER(N{r})),N{r}<=0),"Falta precio",'
          'IF(OR(NOT(ISNUMBER(M{r})),M{r}<0),"Falta costo",'
          'IF(N{r}<M{r},"Precio menor al costo",'
          'IF(AND(LEN(O{r})>0,OR(NOT(ISNUMBER(O{r})),O{r}<=N{r})),"Precio regular debe ser mayor",'
          'IF(NOT(ISNUMBER(P{r})),"Falta stock",'
          'IF(OR(P{r}<0,P{r}<>INT(P{r})),"Falta stock",'
          'IF(AND(Q{r}<>"Sí",Q{r}<>"No"),"Falta disponibilidad","Listo para importar")))))))'),
}
STATUSES = ['Falta precio', 'Falta costo', 'Precio menor al costo', 'Precio regular debe ser mayor', 'Falta stock',
            'Falta disponibilidad', 'Listo para importar']
MONEY, PERCENT, INTEGER = '$ #,##0.00', '0.0 %', '#,##0'

# ── Tipos de inconsistencia, en el orden en que conviene resolverlos ─────────
T_PRICES = 'Mismo nombre y presentación, precios distintos'
T_SAME = 'Mismo nombre y presentación'
T_PACK_PRICE = 'Pack más caro que las unidades sueltas'
T_PACK = 'Pack: precio y stock por pack'
T_SIZE_CONFLICT = 'Tamaño a confirmar'
T_NO_SIZE = 'Presentación sin tamaño'
T_NO_PRESENTATION = 'Sin presentación'
T_NO_BRAND = 'Sin marca'
T_DUP_SKU = 'Código SKU repetido'
T_NO_PRICE = 'Sin precio en el catálogo viejo'
T_PRICE_STATUS = 'Precio sin confirmar en el catálogo viejo'
T_AVAILABLE_NO_STOCK = 'Figuraba disponible sin stock'
T_AVAILABLE_STOCK_MISSING = 'Disponible sin stock cargado'
T_HIDDEN_WITH_STOCK = 'Oculto aunque tenía stock'
T_STATUS = 'Estado raro en el catálogo viejo'
T_CATEGORY_SPELLING = 'Categoría escrita distinto'
T_CATEGORY_LEGACY = 'Categoría de la versión anterior'
T_CATEGORY_UNKNOWN = 'Categoría fuera de la lista'
T_ALCOHOL_CATEGORY = 'Alcohol y categoría no coinciden'
T_ALCOHOL = 'Producto con alcohol'
ISSUE_ORDER = [T_PRICES, T_SAME, T_PACK_PRICE, T_PACK, T_SIZE_CONFLICT, T_NO_SIZE, T_NO_PRESENTATION, T_NO_BRAND,
               T_DUP_SKU, T_NO_PRICE, T_PRICE_STATUS, T_AVAILABLE_NO_STOCK, T_AVAILABLE_STOCK_MISSING,
               T_HIDDEN_WITH_STOCK, T_STATUS, T_CATEGORY_SPELLING, T_CATEGORY_LEGACY, T_CATEGORY_UNKNOWN,
               T_ALCOHOL_CATEGORY, T_ALCOHOL]

# Sugerencia para nombres de categoría de taxonomías anteriores (sólo sugerencia).
LEGACY_SUGGESTION = {'energeticas': 'Energizantes', 'picadas-y-deli': 'Snacks', 'hielo-y-extras': 'Hielo',
                     'jugos-y-saborizadas': 'Jugos', 'complementos': 'Otros', 'vinos-y-espumantes': 'Vinos',
                     'gins-y-vodkas': 'Destilados', 'whisky-y-destilados': 'Destilados'}

# ── Estilos ──────────────────────────────────────────────────────────────────
FONT = 'Calibri'
FILL = {'reference': 'F2F2F2', 'suggestion': 'EDE7F6', 'input': 'FFF2CC', 'formula': 'DDEBF7'}
HEADER_FILL = {'reference': '595959', 'suggestion': '5E3C99', 'input': '9C6500', 'formula': '1F4E79'}
TEXT_COLOR = {'reference': '404040', 'suggestion': '4A2A7A', 'input': '000000', 'formula': '1F4E79'}
THIN = Side(style='thin', color='D0D0D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


# ── Utilidades de texto ──────────────────────────────────────────────────────
def strip_accents(text):
    return ''.join(ch for ch in unicodedata.normalize('NFD', str(text or '')) if not unicodedata.combining(ch))


def norm(text):
    return re.sub(r'[^a-z0-9]+', ' ', strip_accents(text).lower()).strip()


def slugify(text):  # Misma transformación que slugifyCategoryName de store-taxonomy.js
    slug = re.sub(r'[^a-z0-9]+', '-', re.sub(r'[\u0300-\u036f]', '', unicodedata.normalize('NFD', str(text or ''))).lower())
    return slug.strip('-')[:40] or 'otros'


def sort_key(text):
    return norm(text)


def ars(value):
    value = Decimal(value)
    whole = int(value)
    cents = int((value - whole) * 100)
    text = f'{whole:,}'.replace(',', '.')
    return f'${text},{cents:02d}' if cents else f'${text}'


def join_es(items):
    items = list(items)
    if len(items) <= 1:
        return ''.join(items)
    return ', '.join(items[:-1]) + ' y ' + items[-1]


def yes_no(value):
    return 'Sí' if value else 'No'


# ── Fuente: SÓLO data.products ───────────────────────────────────────────────
def read_products(path, expected_sha, accept_different):
    raw = Path(path).read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != expected_sha and not accept_different:
        raise SystemExit(f'El respaldo no coincide con el SHA-256 auditado ({digest}); no se arma la planilla.')
    products = json.loads(raw.decode('utf-8'))['data']['products']
    if not isinstance(products, list) or len(products) != EXPECTED_ROWS:
        raise SystemExit(f'Se esperaban {EXPECTED_ROWS} productos en data.products.')
    for product in products:
        if list(product.keys()) != PRODUCT_KEYS:
            raise SystemExit('Un producto del respaldo tiene columnas inesperadas.')
    # Mismo hash por fila que bootstrap-hosted.mjs (sha256 de JSON.stringify(payload)).
    row_hashes = [hashlib.sha256(json.dumps(p, ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest()
                  for p in products]
    rows_digest = hashlib.sha256('\n'.join(row_hashes).encode('utf-8')).hexdigest()
    return products, digest, rows_digest


# ── Vocabulario de categorías ────────────────────────────────────────────────
def load_taxonomy(repo):
    text = (repo / 'js' / 'core' / 'store-taxonomy.js').read_text(encoding='utf-8')

    def block(marker):
        start = text.index(marker)
        return text[start:text.index(']);', start)]

    def field(body, key):
        match = re.search(r'\b' + key + r":\s*'([^']*)'", body)
        return match.group(1) if match else None

    departments = dict(re.findall(r"id:\s*'([^']+)',\s*name:\s*'([^']+)'", block('export const STORE_DEPARTMENTS')))
    live = []
    for body in re.findall(r'category\(\{(.*?)\}\),', block('export const STORE_CATEGORIES'), re.S):
        synonyms = re.search(r'synonyms:\s*\[([^\]]*)\]', body)
        entry = {'id': field(body, 'id'), 'name': field(body, 'name'),
                 'displayName': field(body, 'displayName') or field(body, 'name'),
                 'department': departments.get(field(body, 'department'), ''),
                 'alcoholic': bool(re.search(r'alcoholic:\s*true', body)),
                 'synonyms': re.findall(r"'([^']*)'", synonyms.group(1)) if synonyms else []}
        if slugify(entry['name']) != entry['id']:
            raise SystemExit(f"La categoría {entry['name']} no slugifica a su id {entry['id']}.")
        live.append(entry)
    legacy = [{'id': field(body, 'id'), 'name': field(body, 'name'), 'alcoholic': bool(re.search(r'alcoholic:\s*true', body))}
              for body in re.findall(r'legacy\(\{(.*?)\}\),', block('export const LEGACY_CATEGORIES'), re.S)]
    migration = next((repo / 'supabase' / 'migrations').glob('20260828120000_*.sql'))
    sql = migration.read_text(encoding='utf-8')
    start = sql.index('add constraint products_verified_canonical_beverage_category check (')
    sql_vocabulary = set(re.findall(r"'([^']+)'", sql[start:sql.index(');', start)]))
    missing = [entry['name'] for entry in live if entry['name'] not in sql_vocabulary]
    if len(live) < 20 or missing:
        raise SystemExit(f'Vocabulario de categorías inconsistente con la base: {missing or len(live)}.')
    return {'live': live, 'legacy': legacy, 'sql_vocabulary': sql_vocabulary, 'migration': migration.name,
            'by_slug': {**{slugify(e['displayName']): e for e in live}, **{e['id']: e for e in live}}}


def suggest_category(product, taxonomy):
    """Devuelve (sugerencia, tipo) con tipo exact | spelling | legacy | inferred | unknown."""
    registered = str(product.get('category') or '').strip()
    live_names = {e['name'] for e in taxonomy['live']}
    if registered in live_names:
        return registered, 'exact'
    slug = slugify(registered)
    if registered and slug in taxonomy['by_slug']:
        return taxonomy['by_slug'][slug]['name'], 'spelling'
    if registered and slug in LEGACY_SUGGESTION:
        return LEGACY_SUGGESTION[slug], 'legacy'
    words = set(norm(f"{product.get('name')} {product.get('description')} {registered}").split())
    matches = {e['name'] for e in taxonomy['live'] for s in e['synonyms'] if s and set(norm(s).split()) <= words}
    if len(matches) == 1:
        return matches.pop(), 'inferred'
    return '', 'unknown'


# ── Tamaño, envase y pack ────────────────────────────────────────────────────
TEXT_SIZE = re.compile(r'(\d+(?:[.,]\d+)?)\s*(ml|cc|lts?|litros?|l|kg|grs?|g)\b', re.I)
SKU_SIZE = re.compile(r'(?:^|-)(\d+)-?(ml|cc|lt|l|kg|g)(?=-|$)')
TEXT_PACK = re.compile(r'\bpack\s*(?:x|de)?\s*(\d+)', re.I)
SKU_PACK = re.compile(r'-pack-x?(\d+)$')
UNIT_FACTOR = {'ml': ('ml', 1), 'cc': ('ml', 1), 'l': ('ml', 1000), 'lt': ('ml', 1000), 'lts': ('ml', 1000),
               'litro': ('ml', 1000), 'litros': ('ml', 1000), 'g': ('g', 1), 'gr': ('g', 1), 'grs': ('g', 1),
               'kg': ('g', 1000)}
SOURCE_PHRASE = {'nombre': 'del nombre', 'presentación': 'de la presentación', 'descripción': 'de la descripción',
                 'código': 'del código'}


def to_amount(number, unit):
    base, factor = UNIT_FACTOR[unit.lower()]
    return base, Decimal(number.replace(',', '.')) * factor


def format_amount(base, amount):
    if base == 'ml' and amount >= 1000:
        liters = (amount / 1000).normalize()
        return f"{format(liters, 'f').replace('.', ',')} L"
    if base == 'g' and amount >= 1000:
        return f"{format((amount / 1000).normalize(), 'f').replace('.', ',')} kg"
    return f"{format(amount.normalize(), 'f')} {base}"


def container_of(*texts):
    for text in texts:
        words = norm(text).split()
        if 'lata' in words or 'latas' in words:
            return 'lata'
        if 'sifon' in words:
            return 'sifón'
        if 'pet' in words:
            return 'botella PET'
        if 'botella' in words:
            return 'botella'
    return ''


def detect_size(product):
    sources = [('nombre', product.get('name')), ('presentación', product.get('presentation')),
               ('descripción', product.get('description'))]
    sizes, packs = {}, {}
    for label, text in sources:
        match = TEXT_SIZE.search(str(text or ''))
        if match:
            sizes[label] = to_amount(match.group(1), match.group(2))
        pack = TEXT_PACK.search(str(text or ''))
        if pack:
            packs[label] = int(pack.group(1))
    sku = str(product.get('sku') or '')
    match = SKU_SIZE.search(sku)
    if match:
        sizes['código'] = to_amount(match.group(1), match.group(2))
    pack = SKU_PACK.search(sku)
    if pack:
        packs['código'] = int(pack.group(1))
    distinct_sizes = sorted(set(sizes.values()))
    distinct_packs = sorted(set(packs.values()))
    size = sizes[next(iter(sizes))] if sizes else None
    units = packs[next(iter(packs))] if packs else None
    container = container_of(product.get('description'), product.get('presentation'), sku.replace('-', ' '))
    text = ''
    if size:
        body = f'{container} {format_amount(*size)}'.strip()
        text = f'Pack x{units} · {body}' if units else body[0].upper() + body[1:]
    return {'size': size, 'pack': units, 'container': container, 'text': text,
            'sources': [label for label in sizes if sizes[label] == size],
            'size_conflict': [(label, format_amount(*value)) for label, value in sizes.items()] if len(distinct_sizes) > 1 else [],
            'pack_conflict': [(label, value) for label, value in packs.items()] if len(distinct_packs) > 1 else [],
            'presentation_has_size': bool(TEXT_SIZE.search(str(product.get('presentation') or '')))}


# ── Análisis ─────────────────────────────────────────────────────────────────
def analyze(products, taxonomy):
    alcoholic_names = {e['name'] for e in taxonomy['live'] if e['alcoholic']} | {e['name'] for e in taxonomy['legacy'] if e['alcoholic']}
    rows = []
    for product in products:
        suggestion, kind = suggest_category(product, taxonomy)
        rows.append({'p': product, 'size': detect_size(product), 'category': suggestion, 'category_kind': kind,
                     'price': Decimal(product['price']) if product.get('price') not in (None, '') else None,
                     'name_key': ' '.join(w for w in norm(product.get('name')).split() if w != 'original'),
                     'issues': []})
    rows.sort(key=lambda row: (sort_key(row['category'] or row['p'].get('category')), sort_key(row['p'].get('name')),
                               row['size']['size'][1] if row['size']['size'] else 0, row['size']['pack'] or 0,
                               row['p']['sku']))
    issues = []

    def add(row, kind, detail, action):
        row['issues'].append(kind)
        issues.append({'sku': row['p']['sku'], 'name': row['p']['name'], 'type': kind, 'detail': detail, 'action': action})

    def label(row):
        text = row['size']['text']
        return (text[0].lower() + text[1:]) if text else 'tamaño sin detectar'

    groups = defaultdict(list)
    for row in rows:
        groups[(norm(row['p'].get('brand')), row['name_key'], norm(row['p'].get('presentation')))].append(row)
    sku_counts = Counter(row['p']['sku'] for row in rows)

    for row in rows:
        p, size = row['p'], row['size']
        presentation = str(p.get('presentation') or '').strip()
        group = groups[(norm(p.get('brand')), row['name_key'], norm(p.get('presentation')))]
        if len(group) > 1:
            names = join_es(sorted({f"«{m['p']['name']}»" for m in group}))
            members = sorted(group, key=lambda m: (m['price'] or 0, m['p']['sku']))
            if len({m['price'] for m in group}) > 1:
                prices = join_es([f"{ars(m['price'])} ({label(m)})" for m in members if m['price'] is not None])
                add(row, T_PRICES,
                    f'En el catálogo viejo hay {len(group)} productos con el mismo nombre y presentación '
                    f'({names} · «{presentation}») y precios distintos: {prices}. Precios vencidos. '
                    'Puede ser que sean tamaños distintos: la presentación no dice el tamaño.',
                    'Confirmá el tamaño de cada uno (columna E) y anotá en Observaciones cómo querés que se vea '
                    f"(por ejemplo «{size['text'] or 'Lata 473 ml'}») para que el cliente no los confunda.")
            else:
                add(row, T_SAME,
                    f'En el catálogo viejo hay {len(group)} productos con el mismo nombre y presentación '
                    f'({names} · «{presentation}») y el mismo precio. Puede ser un producto repetido o tamaños distintos.',
                    'Confirmá si son productos distintos. Si es el mismo, anotalo en Observaciones.')
        if size['pack']:
            singles = [m for m in rows if not m['size']['pack'] and m['size']['size'] == size['size']
                       and norm(m['p'].get('brand')) == norm(p.get('brand')) and m['name_key'] == row['name_key']
                       and m['price'] is not None]
            if singles and row['price'] is not None:
                single = min(singles, key=lambda m: m['p']['sku'])
                total = single['price'] * size['pack']
                if row['price'] > total:
                    add(row, T_PACK_PRICE,
                        f"En el catálogo viejo el pack x{size['pack']} salía {ars(row['price'])} y {size['pack']} "
                        f"unidades sueltas ({single['p']['name']}, {label(single)}) sumaban {ars(total)} "
                        f"({size['pack']} × {ars(single['price'])}). Puede ser un error de carga.",
                        'Revisá que el precio nuevo del pack tenga sentido frente al de la unidad suelta.')
            add(row, T_PACK,
                f"Trae {size['pack']} unidades." + (f" El precio histórico ({ars(row['price'])}, vencido) era por el pack entero."
                                                     if row['price'] is not None else ''),
                'Confirmá si lo vendés como pack. El precio va por el pack entero y el stock se cuenta en packs.')
        if size['size_conflict'] or size['pack_conflict']:
            found = size['size_conflict'] or size['pack_conflict']
            add(row, T_SIZE_CONFLICT,
                'No coincide el tamaño: ' + join_es([f'{SOURCE_PHRASE[k][3:] if SOURCE_PHRASE[k].startswith("de ") else SOURCE_PHRASE[k][4:]} dice {v}'
                                                     for k, v in found]) + '.',
                'Mirá el envase y anotá el tamaño correcto en Observaciones.')
        if not presentation:
            add(row, T_NO_PRESENTATION, 'No tiene presentación registrada.',
                'Anotá en Observaciones cómo se vende (por ejemplo «Lata 473 ml»).')
        elif not size['presentation_has_size']:
            if size['text']:
                origin = join_es([SOURCE_PHRASE[s] for s in size['sources']])
                add(row, T_NO_SIZE,
                    f"La presentación registrada es «{presentation}» y no dice el tamaño. Sugerimos «{size['text']}» "
                    f'(lo sacamos {origin}).',
                    'Confirmá el tamaño mirando el envase. Si la sugerencia está mal, anotalo en Observaciones.')
            else:
                add(row, T_NO_SIZE, f'La presentación registrada es «{presentation}» y no dice el tamaño. No pudimos detectarlo.',
                    'Anotá el tamaño en Observaciones.')
        if not str(p.get('brand') or '').strip():
            add(row, T_NO_BRAND, 'No tiene marca registrada.', 'Anotá la marca en Observaciones.')
        if sku_counts[p['sku']] > 1:
            add(row, T_DUP_SKU, f"El código aparece {sku_counts[p['sku']]} veces en el catálogo viejo.",
                'Avisá antes de importar: el Panel no los puede distinguir.')
        if row['price'] is None or row['price'] <= 0:
            add(row, T_NO_PRICE, 'No tenía precio cargado.', 'Cargá el precio actual.')
        elif p.get('price_status') != 'confirmed':
            add(row, T_PRICE_STATUS, f"El precio histórico figuraba como «{p.get('price_status')}», no confirmado.",
                'Cargá el precio actual.')
        stock = p.get('stock')
        if p.get('available') and stock == 0:
            add(row, T_AVAILABLE_NO_STOCK, 'En el catálogo viejo figuraba disponible pero con stock 0.',
                'Contá el stock actual antes de marcarlo disponible.')
        if p.get('available') and stock is None:
            add(row, T_AVAILABLE_STOCK_MISSING, 'En el catálogo viejo figuraba disponible sin stock cargado.',
                'Contá el stock actual antes de marcarlo disponible.')
        if not p.get('available') and (stock or 0) > 0 and not p.get('is_alcoholic') and p.get('is_active'):
            add(row, T_HIDDEN_WITH_STOCK, f'Tenía stock ({stock}) pero no estaba disponible. Puede que lo hayan ocultado a mano.',
                'Confirmá si lo querés vender.')
        oddities = []
        if not p.get('is_active'):
            oddities.append('dado de baja (inactivo)')
        if not p.get('is_verified'):
            oddities.append('sin verificar')
        if p.get('available') and not (p.get('is_active') and p.get('is_verified')):
            oddities.append('disponible sin estar activo y verificado')
        if p.get('catalog_origin') != 'commercial':
            oddities.append(f"con origen «{p.get('catalog_origin')}» en vez de comercial")
        if oddities:
            add(row, T_STATUS, f'En el catálogo viejo figuraba {join_es(oddities)}.', 'Confirmá si este producto se sigue vendiendo.')
        registered = str(p.get('category') or '').strip()
        if row['category_kind'] == 'spelling':
            add(row, T_CATEGORY_SPELLING, f"La categoría «{registered}» está escrita distinto de la lista de la tienda («{row['category']}»).",
                'Confirmá la categoría sugerida (columna G).')
        elif row['category_kind'] == 'legacy':
            add(row, T_CATEGORY_LEGACY, f"«{registered}» es un nombre de categoría de la versión anterior de la tienda. Sugerimos «{row['category']}».",
                'Confirmá la categoría sugerida o elegí otra de la lista (columna G).')
        elif row['category_kind'] in ('inferred', 'unknown'):
            extra = f" Sugerimos «{row['category']}» por el nombre del producto." if row['category'] else ''
            add(row, T_CATEGORY_UNKNOWN, f"La categoría «{registered or 'vacía'}» no está en la lista de la tienda.{extra}",
                'Elegí una categoría de la lista (columna G).')
        if row['category'] and isinstance(p.get('is_alcoholic'), bool) and p['is_alcoholic'] != (row['category'] in alcoholic_names):
            add(row, T_ALCOHOL_CATEGORY,
                f"La categoría «{row['category']}» {'lleva' if row['category'] in alcoholic_names else 'no lleva'} alcohol, "
                f"pero el producto figuraba {'con' if p['is_alcoholic'] else 'sin'} alcohol.",
                'Confirmá si lleva alcohol y anotalo en Observaciones.')
        if p.get('is_alcoholic'):
            add(row, T_ALCOHOL,
                'Lleva alcohol. Para venderlo online, el comercio necesita la habilitación para vender alcohol '
                'y el control de edad (18+) configurado.',
                'Podés cargar costo, precio y stock. No lo marques disponible hasta tener la habilitación y el '
                'control de edad listos en el Panel.')
    position = {row['p']['sku']: index for index, row in enumerate(rows)}
    issues.sort(key=lambda issue: (ISSUE_ORDER.index(issue['type']), position[issue['sku']]))
    return rows, issues


def counts_for(rows, issues):
    by_type = Counter(issue['type'] for issue in issues)
    products_by_type = {kind: len({i['sku'] for i in issues if i['type'] == kind}) for kind in ISSUE_ORDER}
    return {
        'products': len(rows),
        'presentationWithSize': sum(1 for r in rows if r['size']['presentation_has_size']),
        'presentationWithoutSize': sum(1 for r in rows if not r['size']['presentation_has_size']),
        'suggestedSize': sum(1 for r in rows if r['size']['text']),
        'packs': sum(1 for r in rows if r['size']['pack']),
        'alcoholic': sum(1 for r in rows if r['p'].get('is_alcoholic')),
        'nonAlcoholic': sum(1 for r in rows if not r['p'].get('is_alcoholic')),
        'historicalAvailable': sum(1 for r in rows if r['p'].get('available')),
        'historicalUnavailable': sum(1 for r in rows if not r['p'].get('available')),
        'historicalStockZero': sum(1 for r in rows if r['p'].get('stock') == 0),
        'historicalStockPositive': sum(1 for r in rows if (r['p'].get('stock') or 0) > 0),
        'historicalPriceConfirmed': sum(1 for r in rows if r['p'].get('price_status') == 'confirmed'),
        'historicalPriceExpired': sum(1 for r in rows if r['price'] is not None),
        'categoriesRegistered': len({r['p'].get('category') for r in rows}),
        'categoryKinds': dict(Counter(r['category_kind'] for r in rows)),
        'issueRows': len(issues),
        'productsWithIssues': len({i['sku'] for i in issues}),
        'issueRowsByType': {kind: by_type.get(kind, 0) for kind in ISSUE_ORDER},
        'issueProductsByType': products_by_type,
    }


# ── Planilla ─────────────────────────────────────────────────────────────────
def style_cell(cell, role, *, bold=False, italic=None, wrap=False, align=None, number_format=None, size=10):
    cell.fill = fill(FILL[role])
    cell.font = Font(name=FONT, size=size, bold=bold, italic=(role in ('reference', 'suggestion')) if italic is None else italic,
                     color=TEXT_COLOR[role])
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=wrap, vertical='top', horizontal=align)
    if number_format:
        cell.number_format = number_format


def title(ws, ref, text, size=14):
    ws[ref] = text
    ws[ref].font = Font(name=FONT, size=size, bold=True, color='1F1F1F')


def lines_needed(text, width):
    return max(1, sum(math.ceil(max(1, len(part)) / max(8, width * 1.15)) for part in str(text or '').split('\n')))


def header(ws, row, labels, roles, height=None):
    for column, (text, role) in enumerate(zip(labels, roles), start=1):
        cell = ws.cell(row=row, column=column, value=text)
        cell.fill = fill(HEADER_FILL[role])
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    if height:
        ws.row_dimensions[row].height = height


def build_instructions(ws):
    ws.sheet_view.showGridLines = False
    for letter, width in zip('ABCDEF', (92, 14, 14, 14, 14, 14)):
        ws.column_dimensions[letter].width = width
    title(ws, 'A1', 'La Taba · Catálogo para completar', 16)
    ws['A2'] = 'Qué tenés que hacer'
    ws['A2'].font = Font(name=FONT, size=12, bold=True, color='1F4E79')
    bullets = [
        'En la hoja «Productos» están los 72 productos que recuperamos del catálogo viejo de La Taba. Completá sólo las celdas amarillas.',
        'Por cada producto cargá el costo, el precio de venta, el stock actual y si está disponible para vender (Sí o No).',
        'Escribí los importes como número, sin el signo $ ni puntos de miles: 1500 o 1500,50. El stock va en unidades enteras; en los packs, contá packs.',
        'Una celda vacía quiere decir «no cambiar». Si no sabés un dato, dejala vacía y seguí con el resto.',
        'Las columnas grises son del catálogo de agosto–septiembre de 2026. Ese precio y ese stock están vencidos: sirven sólo de referencia, no los copies sin chequear.',
        'Las columnas lilas («sugerido — confirmar») las armamos leyendo el nombre, el código y la descripción vieja. Si la categoría está mal, elegí otra de la lista; si el tamaño está mal, anotalo en Observaciones.',
        'Margen y recargo no son lo mismo. Si algo te cuesta $1.000 y lo vendés a $1.500, el margen es 33,3 % (de cada venta, un tercio es ganancia) y el recargo es 50 % (le sumaste la mitad al costo). Abajo hay más ejemplos.',
        'Promo: en «Precio de venta» poné lo que cobrás durante la promo, y en «Precio regular antes de la promo» el precio normal, que la tienda muestra tachado. Si no hay promo, dejá vacío el precio regular.',
        'La columna «Estado» te dice qué le falta a cada producto. Cuando dice «Listo para importar», ya tiene todo lo necesario.',
        'En la hoja «Inconsistencias» anotamos cosas raras del catálogo viejo: productos que se ven iguales con precios distintos, packs y productos con alcohol. Revisalas antes de cargar.',
        'Nada se publica solo. Para que un producto aparezca en la tienda tiene que tener precio, stock y foto, y lo tenés que confirmar vos en el Panel. Los productos con alcohol, además, necesitan la habilitación y el control de edad (18+).',
        'Para cargarla: Panel → Catálogo → Importar planilla. Antes de aplicar, el Panel te muestra una vista previa con los cambios: revisala y recién ahí confirmá. No cambies los títulos de las columnas ni los códigos SKU.',
    ]
    row = 3
    for number, text in enumerate(bullets, start=1):
        cell = ws.cell(row=row, column=1, value=f'{number}. {text}')
        cell.font = Font(name=FONT, size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row].height = 15.5 * lines_needed(cell.value, 92) + 3
        row += 1
    row += 1
    ws.cell(row=row, column=1, value='Ejemplos de margen y recargo (se calculan solos)').font = Font(name=FONT, size=12, bold=True, color='1F4E79')
    row += 1
    labels = ['Ejemplo', 'Costo', 'Precio de venta', 'Ganancia por unidad', 'Margen sobre venta', 'Recargo sobre costo']
    header(ws, row, labels, ['reference', 'input', 'input', 'formula', 'formula', 'formula'], 32)
    examples = [
        ('Te cuesta $1.000 y lo vendés a $1.500: de cada venta, un tercio es ganancia (margen 33,3 %) y le sumaste la mitad al costo (recargo 50 %).', 1000, 1500),
        ('Te cuesta $1.000 y lo vendés a $2.000: la mitad de cada venta es ganancia (margen 50 %) y duplicaste el costo (recargo 100 %).', 1000, 2000),
        ('Te cuesta $800 y lo vendés a $1.000: de cada $1.000 que cobrás, $200 son ganancia (margen 20 %); al costo le sumaste un cuarto (recargo 25 %).', 800, 1000),
    ]
    for text, cost, price in examples:
        row += 1
        values = [text, cost, price, f'=C{row}-B{row}', f'=(C{row}-B{row})/C{row}', f'=(C{row}-B{row})/B{row}']
        roles = ['reference', 'input', 'input', 'formula', 'formula', 'formula']
        formats = [None, MONEY, MONEY, MONEY, PERCENT, PERCENT]
        for column, (value, role, number_format) in enumerate(zip(values, roles, formats), start=1):
            style_cell(ws.cell(row=row, column=column, value=value), role, italic=False, wrap=column == 1,
                       number_format=number_format, size=10)
        ws.row_dimensions[row].height = 14 * lines_needed(text, 92) + 4
    row += 2
    ws.cell(row=row, column=1, value='Ejemplo de promo').font = Font(name=FONT, size=12, bold=True, color='1F4E79')
    row += 1
    header(ws, row, ['Ejemplo', 'Precio regular antes de la promo', 'Precio de venta', 'Descuento de la promo'],
           ['reference', 'input', 'input', 'formula'], 44)
    row += 1
    text = 'Normalmente lo vendés a $2.000 y en la promo a $1.600: la tienda muestra $2.000 tachado y cobra $1.600 (descuento 20 %).'
    for column, (value, role, number_format) in enumerate(zip(
            [text, 2000, 1600, f'=(B{row}-C{row})/B{row}'], ['reference', 'input', 'input', 'formula'],
            [None, MONEY, MONEY, PERCENT]), start=1):
        style_cell(ws.cell(row=row, column=column, value=value), role, italic=False, wrap=column == 1, number_format=number_format)
    ws.row_dimensions[row].height = 14 * lines_needed(text, 92) + 4
    row += 2
    ws.cell(row=row, column=1, value='Qué quiere decir cada color').font = Font(name=FONT, size=12, bold=True, color='1F4E79')
    for role, text in [('input', 'Amarillo: lo completás vos.'),
                       ('suggestion', 'Lila: sugerencia nuestra para confirmar. La categoría la podés cambiar con la lista.'),
                       ('reference', 'Gris: dato del catálogo viejo (agosto–septiembre de 2026). Es sólo referencia.'),
                       ('formula', 'Celeste: se calcula solo. No lo escribas.')]:
        row += 1
        style_cell(ws.cell(row=row, column=1, value=text), role, italic=False, size=11)


def build_products(ws, rows):
    ws.sheet_view.showGridLines = False
    title(ws, 'A1', f'Productos para completar ({len(rows)})')
    ws['A2'] = ('Amarillo: completás vos · Lila: sugerencia para confirmar · Gris: catálogo viejo ago–sep 2026 '
                '(vencido) · Celeste: se calcula solo · Celda vacía = no cambiar')
    ws['A2'].font = Font(name=FONT, size=10, italic=True, color='595959')
    header(ws, HEADER_ROW, [c[1] for c in COLUMNS], [c[2] for c in COLUMNS], 58)
    for letter, _, _, width in COLUMNS:
        ws.column_dimensions[letter].width = width
    last = FIRST_ROW + len(rows) - 1
    for offset, row in enumerate(rows):
        r = FIRST_ROW + offset
        p = row['p']
        issues_text = ' · '.join(dict.fromkeys(row['issues']))
        values = {
            'A': p['sku'], 'B': p['name'], 'C': p.get('brand'), 'D': p.get('presentation'), 'E': row['size']['text'] or None,
            'F': p.get('category'), 'G': row['category'] or None, 'H': yes_no(p.get('is_alcoholic')),
            'I': float(row['price']) if row['price'] is not None else None, 'J': p.get('stock'),
            'K': yes_no(p.get('available')), 'L': issues_text or None,
        }
        for letter, _, role, _ in COLUMNS:
            value = FORMULAS[letter].format(r=r) if letter in FORMULAS else values.get(letter)
            cell = ws[f'{letter}{r}']
            cell.value = value
            number_format = {'I': MONEY, 'M': MONEY, 'N': MONEY, 'O': MONEY, 'R': MONEY, 'S': PERCENT, 'T': PERCENT,
                             'U': PERCENT, 'J': INTEGER, 'P': INTEGER}.get(letter)
            align = 'center' if letter in 'HKQ' else None
            style_cell(cell, role, bold=letter in ('A', 'V'), wrap=letter in 'BDELW', align=align, number_format=number_format)
        ws.row_dimensions[r].height = 13.5 * max(lines_needed(issues_text, 42), lines_needed(p['name'], 26),
                                                  lines_needed(row['size']['text'], 25), lines_needed(p.get('presentation'), 21)) + 4
    ws.freeze_panes = f'C{FIRST_ROW}'
    ws.auto_filter.ref = f'A{HEADER_ROW}:W{last}'
    validations = [
        (DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                        errorTitle='Costo', error='Escribí el costo como número mayor o igual a 0, sin el signo $.'), f'M{FIRST_ROW}:M{last}'),
        (DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True,
                        errorTitle='Precio de venta', error='Escribí el precio como número mayor que 0, sin el signo $.'), f'N{FIRST_ROW}:N{last}'),
        (DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                        errorTitle='Precio regular', error='Escribí el precio regular como número, o dejalo vacío si no hay promo.'), f'O{FIRST_ROW}:O{last}'),
        (DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                        errorTitle='Stock', error='Escribí el stock como número entero (0, 1, 2…).'), f'P{FIRST_ROW}:P{last}'),
        (DataValidation(type='list', formula1='"Sí,No"', allow_blank=True,
                        errorTitle='¿Disponible para vender?', error='Elegí Sí o No de la lista.'), f'Q{FIRST_ROW}:Q{last}'),
        (DataValidation(type='list', formula1="'Categorías'!$A$2:$A$" + str(1 + category_count[0]), allow_blank=True,
                        errorTitle='Categoría', error='Elegí una categoría de la lista.'), f'G{FIRST_ROW}:G{last}'),
    ]
    for validation, ref in validations:
        validation.showErrorMessage = True
        validation.add(ref)
        ws.add_data_validation(validation)
    good, bad, neutral = ('C6EFCE', '006100'), ('FFC7CE', '9C0006'), ('FFEB9C', '9C5700')
    status_range = f'V{FIRST_ROW}:V{last}'
    ws.conditional_formatting.add(status_range, CellIsRule(operator='equal', formula=['"Listo para importar"'], stopIfTrue=True,
                                                           fill=fill(good[0]), font=Font(name=FONT, bold=True, color=good[1])))
    ws.conditional_formatting.add(status_range, FormulaRule(
        formula=[f'OR(V{FIRST_ROW}="Precio menor al costo",V{FIRST_ROW}="Precio regular debe ser mayor")'], stopIfTrue=True,
        fill=fill(bad[0]), font=Font(name=FONT, bold=True, color=bad[1])))
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'LEFT(V{FIRST_ROW},5)="Falta"'], stopIfTrue=True,
                                                            fill=fill(neutral[0]), font=Font(name=FONT, bold=True, color=neutral[1])))
    ws.conditional_formatting.add(f'R{FIRST_ROW}:R{last}', CellIsRule(operator='lessThan', formula=['0'],
                                                                       font=Font(name=FONT, bold=True, color=bad[1])))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{HEADER_ROW}:{HEADER_ROW}'


category_count = [0]


def build_issues(ws, issues):
    ws.sheet_view.showGridLines = False
    title(ws, 'A1', f'Cosas para revisar antes de cargar ({len(issues)})')
    ws['A2'] = ('Las encontramos comparando los datos del catálogo viejo. «Puede ser» no es seguro: '
                'confirmalo mirando el producto. Precios y stock son de agosto–septiembre de 2026 (vencidos).')
    ws['A2'].font = Font(name=FONT, size=10, italic=True, color='595959')
    labels = ['Código SKU', 'Producto', 'Tipo', 'Detalle', 'Qué hacer']
    widths = [40, 26, 28, 70, 52]
    header(ws, HEADER_ROW, labels, ['reference', 'reference', 'reference', 'reference', 'input'], 30)
    for letter, width in zip('ABCDE', widths):
        ws.column_dimensions[letter].width = width
    for offset, issue in enumerate(issues):
        r = FIRST_ROW + offset
        for column, key in enumerate(['sku', 'name', 'type', 'detail', 'action'], start=1):
            role = 'input' if key == 'action' else 'reference'
            style_cell(ws.cell(row=r, column=column, value=issue[key]), role, italic=False, wrap=True, bold=key == 'type')
        ws.row_dimensions[r].height = 13.5 * max(lines_needed(issue['detail'], 70), lines_needed(issue['action'], 52),
                                                  lines_needed(issue['type'], 28), lines_needed(issue['name'], 26)) + 4
    ws.freeze_panes = f'A{FIRST_ROW}'
    if issues:
        ws.auto_filter.ref = f'A{HEADER_ROW}:E{FIRST_ROW + len(issues) - 1}'


def build_categories(ws, taxonomy, last_product_row):
    ws.sheet_view.showGridLines = False
    labels = ['Categoría', 'Se muestra en la tienda como', 'Rubro', '¿Lleva alcohol?', 'Cómo la busca la gente',
              'Productos de esta planilla']
    header(ws, 1, labels, ['reference'] * 5 + ['formula'], 32)
    for letter, width in zip('ABCDEF', (20, 26, 26, 12, 52, 14)):
        ws.column_dimensions[letter].width = width
    for offset, entry in enumerate(taxonomy['live']):
        r = 2 + offset
        values = [entry['name'], entry['displayName'], entry['department'], 'Sí (18+)' if entry['alcoholic'] else 'No',
                  ', '.join(entry['synonyms']), f'=COUNTIF(Productos!$G${FIRST_ROW}:$G${last_product_row},A{r})']
        for column, value in enumerate(values, start=1):
            role = 'formula' if column == 6 else 'reference'
            style_cell(ws.cell(row=r, column=column, value=value), role, italic=False, bold=column == 1, wrap=column == 5,
                       align='center' if column in (4, 6) else None)
    note = 2 + len(taxonomy['live']) + 1
    ws.cell(row=note, column=1, value='Usá exactamente estos nombres: son los que acepta la tienda.').font = Font(name=FONT, size=10, italic=True, color='595959')
    ws.freeze_panes = 'A2'


def build_summary(ws, counts, last_product_row):
    ws.sheet_view.showGridLines = False
    for letter, width in zip('ABC', (58, 14, 14)):
        ws.column_dimensions[letter].width = width
    title(ws, 'A1', 'Resumen', 16)
    row = 2

    def section(text, value_label='Cantidad'):
        nonlocal row
        row += 1
        header(ws, row, [text, value_label], ['reference', 'reference'], 22)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='left', vertical='center')

    def line(text, value, role='reference', number_format=INTEGER):
        nonlocal row
        row += 1
        style_cell(ws.cell(row=row, column=1, value=text), 'reference', italic=False)
        style_cell(ws.cell(row=row, column=2, value=value), role, italic=False, align='right', number_format=number_format)

    section('Catálogo recuperado (agosto–septiembre de 2026)')
    for text, key in [('Productos recuperados', 'products'), ('Con el tamaño escrito en la presentación', 'presentationWithSize'),
                      ('Sin el tamaño en la presentación', 'presentationWithoutSize'), ('Con tamaño sugerido para confirmar', 'suggestedSize'),
                      ('Packs (traen más de una unidad)', 'packs'), ('Con alcohol (necesitan habilitación y control de edad)', 'alcoholic'),
                      ('Sin alcohol', 'nonAlcoholic'), ('Estaban disponibles para vender', 'historicalAvailable'),
                      ('No estaban disponibles', 'historicalUnavailable'), ('Tenían stock 0', 'historicalStockZero'),
                      ('Tenían precio cargado (hoy vencido)', 'historicalPriceExpired'), ('Categorías distintas', 'categoriesRegistered'),
                      ('Cosas para revisar (hoja «Inconsistencias»)', 'issueRows'), ('Productos con al menos una cosa para revisar', 'productsWithIssues')]:
        line(text, counts[key])
    section('Cosas para revisar, por tipo', 'Filas')
    for kind in ISSUE_ORDER:
        if counts['issueRowsByType'][kind]:
            line(kind, counts['issueRowsByType'][kind])
    section('Lo que vas completando (se actualiza solo)')
    rng = lambda letter: f'Productos!{letter}{FIRST_ROW}:{letter}{last_product_row}'
    for text, formula, number_format in [
        ('Con precio de venta', f'=COUNT({rng("N")})', INTEGER), ('Con costo', f'=COUNT({rng("M")})', INTEGER),
        ('Con stock actual', f'=COUNT({rng("P")})', INTEGER), ('En promo (con precio regular)', f'=COUNT({rng("O")})', INTEGER),
        ('Marcados «Sí» para vender', f'=COUNTIF({rng("Q")},"Sí")', INTEGER), ('Marcados «No»', f'=COUNTIF({rng("Q")},"No")', INTEGER),
        ('Disponibilidad pendiente', f'=COUNTBLANK({rng("Q")})', INTEGER),
        ('Con observaciones', f'=COUNTA({rng("W")})', INTEGER),
        ('Margen promedio sobre venta (de los que tienen precio y costo)', f'=IFERROR(AVERAGE({rng("S")}),"")', PERCENT),
    ]:
        line(text, formula, 'formula', number_format)
    section('Estado de los productos (se actualiza solo)')
    first_status = row + 1
    for status in STATUSES:
        line(status, f'=COUNTIF({rng("V")},"{status}")', 'formula')
    line('Total (tiene que dar 72)', f'=SUM(B{first_status}:B{row})', 'formula')
    line('Avance: listos para importar sobre el total', f'=COUNTIF({rng("V")},"Listo para importar")/{last_product_row - FIRST_ROW + 1}',
         'formula', PERCENT)


def build_workbook(rows, issues, counts, taxonomy, meta):
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in SHEETS}
    category_count[0] = len(taxonomy['live'])
    last_product_row = FIRST_ROW + len(rows) - 1
    build_instructions(sheets['Instrucciones'])
    build_products(sheets['Productos'], rows)
    build_issues(sheets['Inconsistencias'], issues)
    build_categories(sheets['Categorías'], taxonomy, last_product_row)
    build_summary(sheets['Resumen'], counts, last_product_row)
    sheets['Instrucciones'].sheet_properties.tabColor = '1F4E79'
    sheets['Productos'].sheet_properties.tabColor = 'BF8F00'
    sheets['Inconsistencias'].sheet_properties.tabColor = 'C00000'
    workbook.active = SHEETS.index('Instrucciones')
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    workbook.properties.title = 'La Taba · Catálogo para completar'
    workbook.properties.creator = 'La Taba'
    workbook.properties.description = (f"72 productos recuperados del respaldo parcial del 18/09/2026 (SHA-256 {meta['sha256']}). "
                                       'Sin costos, precios actuales ni stock inventados.')
    return workbook


# ── Verificación del archivo guardado ────────────────────────────────────────
CELL_REF = re.compile(r'(?<![A-Za-z$!])\$?([A-W])\$?(\d+)')


def verify_workbook(path, rows, taxonomy):
    workbook = load_workbook(path)
    checks = {}
    assert workbook.sheetnames == SHEETS, workbook.sheetnames
    ws = workbook['Productos']
    header_rows = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == 'SKU']
    assert header_rows and header_rows[0] == HEADER_ROW, header_rows
    assert [ws.cell(row=HEADER_ROW, column=i + 1).value for i in range(len(COLUMNS))] == [c[1] for c in COLUMNS]
    assert ws.cell(row=HEADER_ROW, column=len(COLUMNS) + 1).value is None
    last = FIRST_ROW + len(rows) - 1
    skus = [ws[f'A{r}'].value for r in range(FIRST_ROW, last + 1)]
    assert len(set(skus)) == EXPECTED_ROWS == len(skus) and set(skus) == {row['p']['sku'] for row in rows}
    assert ws[f'A{last + 1}'].value is None
    for r in range(FIRST_ROW, last + 1):
        for letter in INPUT_LETTERS:
            assert ws[f'{letter}{r}'].value is None, f'{letter}{r} debería estar vacía'
        for letter, template in FORMULAS.items():
            formula = ws[f'{letter}{r}'].value
            assert formula == template.format(r=r), f'{letter}{r}'
            assert {int(n) for _, n in CELL_REF.findall(formula)} == {r}, f'{letter}{r} apunta a otra fila'
        assert isinstance(ws[f'I{r}'].value, (int, float)) and ws[f'I{r}'].number_format == MONEY
        assert all(ws[f'{x}{r}'].number_format == MONEY for x in 'MNOR')
        assert all(ws[f'{x}{r}'].number_format == PERCENT for x in 'STU')
    checks['formulaRowsVerified'] = len(rows)
    ranges = {str(v.sqref): (v.type, v.operator, v.formula1) for v in ws.data_validations.dataValidation}
    expected = {f'M{FIRST_ROW}:M{last}': ('decimal', 'greaterThanOrEqual', '0'), f'N{FIRST_ROW}:N{last}': ('decimal', 'greaterThan', '0'),
                f'O{FIRST_ROW}:O{last}': ('decimal', 'greaterThanOrEqual', '0'), f'P{FIRST_ROW}:P{last}': ('whole', 'greaterThanOrEqual', '0'),
                f'Q{FIRST_ROW}:Q{last}': ('list', None, '"Sí,No"'),
                f'G{FIRST_ROW}:G{last}': ('list', None, f"'Categorías'!$A$2:$A${1 + len(taxonomy['live'])}")}
    assert ranges == expected, ranges
    checks['dataValidations'] = sorted(ranges)
    assert ws.freeze_panes == f'C{FIRST_ROW}'
    for name in ('Instrucciones', 'Inconsistencias', 'Categorías', 'Resumen'):
        other = workbook[name]
        assert all(other.cell(row=r, column=1).value != 'SKU' for r in range(1, other.max_row + 1)), name
    categories = workbook['Categorías']
    names = [categories[f'A{r}'].value for r in range(2, 2 + len(taxonomy['live']))]
    assert names == [e['name'] for e in taxonomy['live']]
    assert all(row['category'] in names for row in rows)
    summary = workbook['Resumen']
    formulas = [c.value for col in summary.iter_cols(min_col=2, max_col=2) for c in col if isinstance(c.value, str) and c.value.startswith('=')]
    for formula in formulas:
        for _, number in re.findall(r'Productos!([A-Z])(\d+)', formula):
            assert int(number) in (FIRST_ROW, last), formula
    checks['summaryFormulas'] = len(formulas)
    checks['reopenedWithOpenpyxl'] = True
    return checks


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    parser.add_argument('--backup', required=True)
    parser.add_argument('--out', required=True)
    parser.add_argument('--summary')
    parser.add_argument('--repo', default=str(Path(__file__).resolve().parents[2]))
    parser.add_argument('--expected-sha256', default=EXPECTED_SOURCE_SHA256)
    parser.add_argument('--accept-different-backup', action='store_true')
    args = parser.parse_args(argv)
    repo = Path(args.repo).resolve()
    out = Path(args.out).resolve()
    summary_path = Path(args.summary).resolve() if args.summary else out.with_name(out.stem + '.summary.json')
    for target in (out, summary_path):
        if target.is_relative_to(repo):
            raise SystemExit('La planilla y su resumen van fuera del repositorio: el repositorio es público.')
    products, digest, rows_digest = read_products(args.backup, args.expected_sha256, args.accept_different_backup)
    taxonomy = load_taxonomy(repo)
    rows, issues = analyze(products, taxonomy)
    counts = counts_for(rows, issues)
    workbook = build_workbook(rows, issues, counts, taxonomy, {'sha256': digest})
    out.parent.mkdir(parents=True, exist_ok=True)
    temporary = out.with_name(out.stem + '.tmp.xlsx')
    workbook.save(temporary)
    checks = verify_workbook(temporary, rows, taxonomy)
    os.replace(temporary, out)
    workbook_sha = hashlib.sha256(out.read_bytes()).hexdigest()
    examples = {kind: next(({'sku': i['sku'], 'detail': i['detail']} for i in issues if i['type'] == kind), None)
                for kind in (T_PRICES, T_PACK_PRICE, T_NO_SIZE, T_PACK, T_ALCOHOL)}
    summary = {
        'schemaVersion': 1,
        'generatedAt': dt.datetime.now(dt.timezone.utc).isoformat(timespec='seconds'),
        'generator': 'scripts/catalog/build-walter-catalog-sheet.py',
        'workbook': {'file': out.name, 'sha256': workbook_sha, 'bytes': out.stat().st_size, 'sheets': SHEETS},
        'source': {'file': Path(args.backup).name, 'sha256': digest, 'expectedSha256': args.expected_sha256,
                   'matchesAuditedSha256': digest == EXPECTED_SOURCE_SHA256, 'readKeys': ['data.products'],
                   'rows': len(products), 'rowsSha256Digest': rows_digest,
                   'rowsSha256DigestMethod': 'sha256 of the per-row sha256(JSON.stringify(payload)) joined by \\n in backup order; '
                                             'stage-historical-catalog-drafts.mjs prints the same sourceRowsDigest'},
        'rules': {'costsInvented': False, 'currentPricesInvented': False, 'stockInvented': False,
                  'historicalLabel': HISTORIC, 'suggestionLabel': SUGGESTED, 'inputsPrefilled': False},
        'columnContract': {
            'sheet': 'Productos', 'headerRow': HEADER_ROW, 'firstDataRow': FIRST_ROW, 'lastDataRow': FIRST_ROW + len(rows) - 1,
            'headerDetection': 'first row whose first cell is exactly "SKU"',
            'columns': [{'letter': l, 'header': h, 'role': role, **({'formula': FORMULAS[l].format(r=FIRST_ROW)} if l in FORMULAS else {})}
                        for l, h, role, _ in COLUMNS],
            'inputColumns': INPUT_LETTERS, 'statusValues': STATUSES, 'blankMeans': 'no cambiar',
            'numberFormats': {'money': MONEY, 'percent': PERCENT},
            'validations': {'M': 'decimal >= 0', 'N': 'decimal > 0', 'O': 'decimal >= 0', 'P': 'whole >= 0', 'Q': 'list Sí,No',
                            'G': "list 'Categorías'!A2:A" + str(1 + len(taxonomy['live']))}},
        'categories': {'source': 'js/core/store-taxonomy.js', 'databaseCheck': taxonomy['migration'],
                       'live': [e['name'] for e in taxonomy['live']],
                       'registered': dict(Counter(r['p'].get('category') for r in rows))},
        'counts': counts,
        'issueTypes': [{'type': kind, 'rows': counts['issueRowsByType'][kind], 'products': counts['issueProductsByType'][kind]}
                       for kind in ISSUE_ORDER],
        'issueExamples': examples,
        'verification': checks,
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'workbook': str(out), 'summary': str(summary_path), 'sourceSha256Matches': digest == EXPECTED_SOURCE_SHA256,
                      'products': counts['products'], 'issueRows': counts['issueRows'],
                      'issueRowsByType': {k: v for k, v in counts['issueRowsByType'].items() if v},
                      'verification': checks}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
